#!/usr/bin/env python3
# -*- coding: utf-8 -*-

"""
Incorporación de metadatos enriquecidos al RIS maestro
Revisión sistemática PRISMA / PRISMA-S

Funciones principales:
1. Lee el RIS maestro consolidado y deduplicado.
2. Lee el CSV enriquecido mediante OpenAlex, Crossref
   y Springer Nature API.
3. Relaciona ambos conjuntos mediante el identificador
   interno único de cada referencia.
4. Completa únicamente los metadatos ausentes o añade
   información complementaria sin eliminar los datos
   bibliográficos originales.
5. Incorpora:
   - abstracts recuperados;
   - palabras clave;
   - DOI ausentes;
   - autores ausentes;
   - año ausente;
   - idioma ausente;
   - URL recuperadas;
   - fuente de publicación ausente.
6. Lee opcionalmente el Excel de revisión manual de los
   registros pendientes.
7. Incorpora del Excel únicamente metadatos bibliográficos
   verificables, pero NO utiliza las paráfrasis manuales
   como si fueran abstracts originales.
8. Conserva la trazabilidad de la fuente de enriquecimiento.
9. Genera un nuevo RIS enriquecido sin modificar
   el RIS maestro original.
10. Genera un informe CSV de las modificaciones realizadas.
"""

from pathlib import Path
from collections import defaultdict
import csv
import re

from openpyxl import load_workbook


# ==========================================================
# CONFIGURACIÓN
# ==========================================================

ARCHIVO_RIS_MAESTRO = Path(
    "ris_normalizados/TFM_bibliografia_consolidada.ris"
)

ARCHIVO_CSV_ENRIQUECIDO = Path(
    "ris_normalizados/"
    "TFM_registros_sin_abstract_enriquecidos_springer.csv"
)

ARCHIVO_EXCEL_REVISION_MANUAL = Path(
    "ris_normalizados/"
    "TFM_12_registros_pendientes_revision_manual_COMPLETADO.xlsx"
)

ARCHIVO_RIS_SALIDA = Path(
    "ris_normalizados/"
    "TFM_bibliografia_consolidada_ENRIQUECIDA.ris"
)

ARCHIVO_INFORME = Path(
    "ris_normalizados/"
    "TFM_informe_incorporacion_metadatos.csv"
)

NUMERO_ESPERADO_REGISTROS = 5479


# ==========================================================
# LECTURA DE RIS
# ==========================================================

def leer_ris(ruta):
    """
    Lee un fichero RIS y devuelve una lista de registros.

    Cada registro se representa mediante un diccionario:
        etiqueta RIS -> lista de valores
    """

    registros = []
    registro_actual = defaultdict(list)
    ultima_etiqueta = None

    with ruta.open(
        "r",
        encoding="utf-8-sig",
        errors="replace"
    ) as archivo:

        for linea_bruta in archivo:

            linea = linea_bruta.rstrip("\r\n")

            coincidencia = re.match(
                r"^([A-Z0-9]{2})  - ?(.*)$",
                linea
            )

            if coincidencia:

                etiqueta = coincidencia.group(1)
                valor = coincidencia.group(2).strip()

                if etiqueta == "TY":

                    if registro_actual:
                        registros.append(
                            dict(registro_actual)
                        )

                    registro_actual = defaultdict(list)

                registro_actual[etiqueta].append(
                    valor
                )

                ultima_etiqueta = etiqueta

                if etiqueta == "ER":

                    registros.append(
                        dict(registro_actual)
                    )

                    registro_actual = defaultdict(list)
                    ultima_etiqueta = None

            elif ultima_etiqueta and linea.strip():

                registro_actual[
                    ultima_etiqueta
                ][-1] += (
                    " " + linea.strip()
                )

    if registro_actual:
        registros.append(
            dict(registro_actual)
        )

    return registros


# ==========================================================
# FUNCIONES AUXILIARES
# ==========================================================

def obtener_primer_valor(
    registro,
    *etiquetas
):
    """
    Devuelve el primer valor no vacío encontrado
    entre las etiquetas indicadas.
    """

    for etiqueta in etiquetas:

        valores = registro.get(
            etiqueta,
            []
        )

        if valores:

            valor = str(
                valores[0]
            ).strip()

            if valor:
                return valor

    return ""


def limpiar_espacios(texto):
    """
    Elimina espacios redundantes.
    """

    if not texto:
        return ""

    return re.sub(
        r"\s+",
        " ",
        str(texto)
    ).strip()


def normalizar_doi(doi):
    """
    Convierte el DOI a un formato uniforme.
    """

    if not doi:
        return ""

    doi = str(doi).strip().lower()

    doi = re.sub(
        r"^https?://(?:dx\.)?doi\.org/",
        "",
        doi
    )

    doi = re.sub(
        r"^doi:\s*",
        "",
        doi
    )

    return doi.rstrip(
        " .;,)]}"
    )


def separar_valores(texto):
    """
    Divide una cadena de valores separados mediante
    punto y coma.

    Se utiliza fundamentalmente para palabras clave
    y autores procedentes del CSV enriquecido.
    """

    if not texto:
        return []

    return [
        limpiar_espacios(valor)
        for valor in str(texto).split(";")
        if limpiar_espacios(valor)
    ]


def anadir_valor_sin_duplicar(
    registro,
    etiqueta,
    valor
):
    """
    Añade un valor a una etiqueta RIS evitando duplicados.
    """

    valor = limpiar_espacios(valor)

    if not valor:
        return False

    if etiqueta not in registro:
        registro[etiqueta] = []

    valores_existentes = {
        limpiar_espacios(elemento).lower()
        for elemento in registro[etiqueta]
        if limpiar_espacios(elemento)
    }

    if valor.lower() in valores_existentes:
        return False

    registro[etiqueta].append(
        valor
    )

    return True


# ==========================================================
# LECTURA DEL CSV ENRIQUECIDO
# ==========================================================

def leer_csv_enriquecido(ruta):
    """
    Lee el CSV enriquecido y lo indexa mediante
    identificador interno.
    """

    datos = {}

    with ruta.open(
        "r",
        encoding="utf-8-sig",
        newline=""
    ) as archivo:

        lector = csv.DictReader(
            archivo
        )

        for fila in lector:

            identificador = (
                fila.get(
                    "identificador_interno"
                )
                or ""
            ).strip()

            if identificador:

                if identificador in datos:

                    raise ValueError(
                        "Identificador duplicado en CSV: "
                        f"{identificador}"
                    )

                datos[identificador] = fila

    return datos


# ==========================================================
# LECTURA DE LA REVISIÓN MANUAL
# ==========================================================

def leer_revision_manual(ruta):
    """
    Lee el Excel de revisión manual.

    Solo se utilizan posteriormente metadatos
    bibliográficos verificables.

    Las paráfrasis del contenido NO se incorporan
    como abstract bibliográfico.
    """

    if not ruta.exists():
        return {}

    libro = load_workbook(
        ruta,
        data_only=True
    )

    hoja = libro["12 pendientes"]

    encabezados = [
        celda.value
        for celda in hoja[1]
    ]

    datos = {}

    for valores in hoja.iter_rows(
        min_row=2,
        values_only=True
    ):

        fila = dict(
            zip(
                encabezados,
                valores
            )
        )

        identificador = str(
            fila.get(
                "Identificador interno"
            )
            or ""
        ).strip()

        if identificador:
            datos[identificador] = fila

    return datos


# ==========================================================
# ELECCIÓN DE METADATOS RECUPERADOS
# ==========================================================

def obtener_autores_recuperados(fila):
    """
    Selecciona los autores recuperados utilizando
    primero Springer, después Crossref y finalmente
    OpenAlex.

    Solo se utilizarán si el RIS carece de autores.
    """

    for campo in (
        "autores_springer",
        "autores_crossref",
        "autores_openalex"
    ):

        valor = (
            fila.get(campo)
            or ""
        ).strip()

        if valor:
            return separar_valores(
                valor
            )

    return []


def obtener_anio_recuperado(fila):
    """
    Selecciona un año recuperado.
    """

    for campo in (
        "anio_springer",
        "anio_crossref",
        "anio_openalex"
    ):

        valor = str(
            fila.get(campo)
            or ""
        ).strip()

        coincidencia = re.search(
            r"\b(19|20)\d{2}\b",
            valor
        )

        if coincidencia:
            return coincidencia.group(0)

    return ""


def obtener_publicacion_recuperada(fila):
    """
    Selecciona el nombre de la publicación recuperada.
    """

    for campo in (
        "fuente_publicacion_springer",
        "fuente_publicacion_crossref",
        "fuente_publicacion_openalex"
    ):

        valor = (
            fila.get(campo)
            or ""
        ).strip()

        if valor:
            return limpiar_espacios(
                valor
            )

    return ""


def obtener_idioma_recuperado(fila):
    """
    Obtiene y homogeneiza el idioma recuperado.
    """

    for campo in (
        "idioma_springer",
        "idioma_openalex"
    ):

        valor = (
            fila.get(campo)
            or ""
        ).strip().lower()

        if valor:

            if valor in {
                "en",
                "eng",
                "english"
            }:
                return "English"

            return valor

    return ""


def obtener_urls_recuperadas(fila):
    """
    Recupera las URL encontradas en las distintas fuentes.
    """

    urls = []

    for campo in (
        "url_springer",
        "url_crossref",
        "url_openalex"
    ):

        valor = (
            fila.get(campo)
            or ""
        ).strip()

        if valor and valor not in urls:
            urls.append(valor)

    return urls


def obtener_doi_recuperado(fila):
    """
    Selecciona un DOI recuperado de forma automática.
    """

    for campo in (
        "doi_recuperado",
        "doi_springer",
        "doi_crossref",
        "doi_openalex"
    ):

        doi = normalizar_doi(
            fila.get(campo)
            or ""
        )

        if doi:
            return doi

    return ""


# ==========================================================
# INCORPORACIÓN DEL ENRIQUECIMIENTO AUTOMÁTICO
# ==========================================================

def enriquecer_registro(
    registro,
    fila,
    informe
):
    """
    Incorpora al registro RIS los metadatos recuperados.

    Los metadatos originales válidos no se sustituyen.
    """

    identificador = obtener_primer_valor(
        registro,
        "ID"
    )

    # ------------------------------------------------------
    # ABSTRACT
    # ------------------------------------------------------

    abstract_existente = obtener_primer_valor(
        registro,
        "AB",
        "N2"
    )

    abstract_recuperado = limpiar_espacios(
        fila.get(
            "abstract_recuperado"
        )
        or ""
    )

    if (
        not abstract_existente
        and abstract_recuperado
    ):

        registro["AB"] = [
            abstract_recuperado
        ]

        fuente_abstract = (
            fila.get(
                "fuente_abstract_recuperado"
            )
            or ""
        ).strip()

        informe.append({
            "identificador": identificador,
            "campo": "AB",
            "accion": "COMPLETADO",
            "fuente": fuente_abstract,
            "valor": "Abstract incorporado"
        })

        if fuente_abstract:

            anadir_valor_sin_duplicar(
                registro,
                "N1",
                (
                    "Abstract enriquecido desde: "
                    f"{fuente_abstract}"
                )
            )

    # ------------------------------------------------------
    # PALABRAS CLAVE
    # ------------------------------------------------------

    palabras_clave = []

    for campo in (
        "palabras_clave_recuperadas",
        "palabras_clave_springer"
    ):

        palabras_clave.extend(
            separar_valores(
                fila.get(campo)
                or ""
            )
        )

    for palabra_clave in palabras_clave:

        incorporada = (
            anadir_valor_sin_duplicar(
                registro,
                "KW",
                palabra_clave
            )
        )

        if incorporada:

            informe.append({
                "identificador": identificador,
                "campo": "KW",
                "accion": "AÑADIDO",
                "fuente":
                    "OpenAlex/Springer Nature",
                "valor": palabra_clave
            })

    # ------------------------------------------------------
    # DOI
    # ------------------------------------------------------

    doi_existente = normalizar_doi(
        obtener_primer_valor(
            registro,
            "DO"
        )
    )

    if not doi_existente:

        doi_recuperado = (
            obtener_doi_recuperado(
                fila
            )
        )

        if doi_recuperado:

            registro["DO"] = [
                doi_recuperado
            ]

            informe.append({
                "identificador": identificador,
                "campo": "DO",
                "accion": "COMPLETADO",
                "fuente":
                    "Enriquecimiento automático",
                "valor": doi_recuperado
            })

    # ------------------------------------------------------
    # AUTORES
    # ------------------------------------------------------

    autores_existentes = (
        registro.get("AU")
        or registro.get("A1")
        or []
    )

    if not any(
        limpiar_espacios(autor)
        for autor in autores_existentes
    ):

        autores_recuperados = (
            obtener_autores_recuperados(
                fila
            )
        )

        if autores_recuperados:

            registro["AU"] = (
                autores_recuperados
            )

            informe.append({
                "identificador": identificador,
                "campo": "AU",
                "accion": "COMPLETADO",
                "fuente":
                    "Springer/Crossref/OpenAlex",
                "valor":
                    "; ".join(
                        autores_recuperados
                    )
            })

    # ------------------------------------------------------
    # AÑO
    # ------------------------------------------------------

    anio_existente = obtener_primer_valor(
        registro,
        "PY",
        "Y1",
        "DA"
    )

    if not anio_existente:

        anio_recuperado = (
            obtener_anio_recuperado(
                fila
            )
        )

        if anio_recuperado:

            registro["PY"] = [
                anio_recuperado
            ]

            informe.append({
                "identificador": identificador,
                "campo": "PY",
                "accion": "COMPLETADO",
                "fuente":
                    "Springer/Crossref/OpenAlex",
                "valor": anio_recuperado
            })

    # ------------------------------------------------------
    # IDIOMA
    # ------------------------------------------------------

    idioma_existente = obtener_primer_valor(
        registro,
        "LA"
    )

    if not idioma_existente:

        idioma_recuperado = (
            obtener_idioma_recuperado(
                fila
            )
        )

        if idioma_recuperado:

            registro["LA"] = [
                idioma_recuperado
            ]

            informe.append({
                "identificador": identificador,
                "campo": "LA",
                "accion": "COMPLETADO",
                "fuente":
                    "Springer/OpenAlex",
                "valor": idioma_recuperado
            })

    # ------------------------------------------------------
    # PUBLICACIÓN
    # ------------------------------------------------------

    publicacion_existente = (
        obtener_primer_valor(
            registro,
            "JO",
            "JF",
            "JA",
            "T2"
        )
    )

    if not publicacion_existente:

        publicacion_recuperada = (
            obtener_publicacion_recuperada(
                fila
            )
        )

        if publicacion_recuperada:

            registro["JO"] = [
                publicacion_recuperada
            ]

            informe.append({
                "identificador": identificador,
                "campo": "JO",
                "accion": "COMPLETADO",
                "fuente":
                    "Springer/Crossref/OpenAlex",
                "valor":
                    publicacion_recuperada
            })

    # ------------------------------------------------------
    # URL
    # ------------------------------------------------------

    for url in obtener_urls_recuperadas(
        fila
    ):

        incorporada = (
            anadir_valor_sin_duplicar(
                registro,
                "UR",
                url
            )
        )

        if incorporada:

            informe.append({
                "identificador": identificador,
                "campo": "UR",
                "accion": "AÑADIDO",
                "fuente":
                    "Enriquecimiento automático",
                "valor": url
            })


# ==========================================================
# INCORPORACIÓN DE LA REVISIÓN MANUAL
# ==========================================================

def incorporar_revision_manual(
    registro,
    fila_manual,
    informe
):
    """
    Incorpora únicamente metadatos bibliográficos
    verificables obtenidos durante la revisión manual.

    Las paráfrasis de contenido NO se incorporan
    al campo AB.
    """

    identificador = obtener_primer_valor(
        registro,
        "ID"
    )

    # ------------------------------------------------------
    # DOI MANUAL
    # ------------------------------------------------------

    doi_existente = normalizar_doi(
        obtener_primer_valor(
            registro,
            "DO"
        )
    )

    doi_manual = normalizar_doi(
        fila_manual.get("DOI")
        or ""
    )

    if (
        not doi_manual
        and fila_manual.get(
            "Enlace DOI"
        )
    ):

        doi_manual = normalizar_doi(
            fila_manual.get(
                "Enlace DOI"
            )
        )

    if (
        not doi_existente
        and doi_manual
    ):

        registro["DO"] = [
            doi_manual
        ]

        informe.append({
            "identificador": identificador,
            "campo": "DO",
            "accion":
                "COMPLETADO MANUALMENTE",
            "fuente":
                "Revisión manual",
            "valor": doi_manual
        })

        anadir_valor_sin_duplicar(
            registro,
            "N1",
            (
                "DOI verificado mediante "
                "revisión manual."
            )
        )

    # ------------------------------------------------------
    # URL DE LA FUENTE MANUAL
    # ------------------------------------------------------

    url_manual = limpiar_espacios(
        fila_manual.get(
            "URL de la fuente manual"
        )
        or ""
    )

    if url_manual:

        incorporada = (
            anadir_valor_sin_duplicar(
                registro,
                "UR",
                url_manual
            )
        )

        if incorporada:

            informe.append({
                "identificador": identificador,
                "campo": "UR",
                "accion":
                    "AÑADIDO MANUALMENTE",
                "fuente":
                    "Revisión manual",
                "valor": url_manual
            })


# ==========================================================
# ESCRITURA DEL RIS
# ==========================================================

ORDEN_RIS = [
    "TY",
    "ID",
    "TI",
    "AU",
    "PY",
    "DA",
    "JO",
    "T2",
    "VL",
    "IS",
    "SP",
    "EP",
    "SN",
    "DO",
    "UR",
    "LA",
    "KW",
    "AB",
    "PB",
    "DB",
    "N1",
]


def escribir_ris(
    registros,
    ruta
):
    """
    Escribe los registros en formato RIS.
    """

    ruta.parent.mkdir(
        parents=True,
        exist_ok=True
    )

    with ruta.open(
        "w",
        encoding="utf-8",
        newline="\n"
    ) as archivo:

        for registro in registros:

            etiquetas_escritas = set()

            for etiqueta in ORDEN_RIS:

                for valor in registro.get(
                    etiqueta,
                    []
                ):

                    if (
                        valor
                        and etiqueta != "ER"
                    ):

                        archivo.write(
                            f"{etiqueta}  - "
                            f"{valor}\n"
                        )

                etiquetas_escritas.add(
                    etiqueta
                )

            # Conserva cualquier campo RIS adicional
            # que estuviera presente en el maestro.
            for etiqueta, valores in (
                registro.items()
            ):

                if (
                    etiqueta
                    in etiquetas_escritas
                    or etiqueta == "ER"
                ):
                    continue

                for valor in valores:

                    if valor:

                        archivo.write(
                            f"{etiqueta}  - "
                            f"{valor}\n"
                        )

            archivo.write(
                "ER  - \n\n"
            )


# ==========================================================
# INFORME DE MODIFICACIONES
# ==========================================================

def escribir_informe(
    informe,
    ruta
):
    """
    Genera un CSV con todas las modificaciones efectuadas.
    """

    campos = [
        "identificador",
        "campo",
        "accion",
        "fuente",
        "valor",
    ]

    with ruta.open(
        "w",
        encoding="utf-8-sig",
        newline=""
    ) as archivo:

        escritor = csv.DictWriter(
            archivo,
            fieldnames=campos
        )

        escritor.writeheader()
        escritor.writerows(
            informe
        )


# ==========================================================
# PROCESO PRINCIPAL
# ==========================================================

def principal():

    print()
    print(
        "INCORPORACIÓN DE METADATOS AL RIS MAESTRO"
    )
    print("=" * 70)

    # ------------------------------------------------------
    # COMPROBACIÓN DE ARCHIVOS
    # ------------------------------------------------------

    if not ARCHIVO_RIS_MAESTRO.exists():

        print(
            "ERROR: No se encuentra el RIS maestro:"
        )

        print(
            ARCHIVO_RIS_MAESTRO.resolve()
        )

        return

    if not ARCHIVO_CSV_ENRIQUECIDO.exists():

        print(
            "ERROR: No se encuentra el CSV enriquecido:"
        )

        print(
            ARCHIVO_CSV_ENRIQUECIDO.resolve()
        )

        return

    # ------------------------------------------------------
    # LECTURA
    # ------------------------------------------------------

    registros = leer_ris(
        ARCHIVO_RIS_MAESTRO
    )

    enriquecimiento = (
        leer_csv_enriquecido(
            ARCHIVO_CSV_ENRIQUECIDO
        )
    )

    revision_manual = (
        leer_revision_manual(
            ARCHIVO_EXCEL_REVISION_MANUAL
        )
    )

    print(
        f"Registros del RIS maestro: "
        f"{len(registros)}"
    )

    print(
        f"Registros del CSV enriquecido: "
        f"{len(enriquecimiento)}"
    )

    print(
        f"Registros de revisión manual: "
        f"{len(revision_manual)}"
    )

    # ------------------------------------------------------
    # CONTROL DEL NÚMERO DE REGISTROS
    # ------------------------------------------------------

    if (
        len(registros)
        != NUMERO_ESPERADO_REGISTROS
    ):

        raise ValueError(
            "El RIS maestro no contiene el número "
            "esperado de registros. "
            f"Esperados: "
            f"{NUMERO_ESPERADO_REGISTROS}. "
            f"Encontrados: {len(registros)}."
        )

    # ------------------------------------------------------
    # COMPROBACIÓN DE IDs
    # ------------------------------------------------------

    indice_ids = {}

    for posicion, registro in enumerate(
        registros
    ):

        identificador = (
            obtener_primer_valor(
                registro,
                "ID"
            )
        )

        if not identificador:

            raise ValueError(
                "Se ha encontrado un registro "
                "sin identificador interno."
            )

        if identificador in indice_ids:

            raise ValueError(
                "Identificador interno duplicado "
                "en RIS maestro: "
                f"{identificador}"
            )

        indice_ids[
            identificador
        ] = posicion

    # ------------------------------------------------------
    # INCORPORACIÓN DE DATOS
    # ------------------------------------------------------

    informe = []

    registros_enriquecidos = 0
    registros_manual = 0

    for identificador, fila in (
        enriquecimiento.items()
    ):

        if identificador not in indice_ids:

            print(
                "ADVERTENCIA: ID del CSV no localizado "
                f"en RIS: {identificador}"
            )

            continue

        posicion = indice_ids[
            identificador
        ]

        numero_modificaciones_antes = len(
            informe
        )

        enriquecer_registro(
            registros[posicion],
            fila,
            informe
        )

        if (
            len(informe)
            > numero_modificaciones_antes
        ):

            registros_enriquecidos += 1

    # ------------------------------------------------------
    # INCORPORACIÓN DE REVISIÓN MANUAL
    # ------------------------------------------------------

    for identificador, fila_manual in (
        revision_manual.items()
    ):

        if identificador not in indice_ids:

            print(
                "ADVERTENCIA: ID manual no localizado "
                f"en RIS: {identificador}"
            )

            continue

        posicion = indice_ids[
            identificador
        ]

        numero_modificaciones_antes = len(
            informe
        )

        incorporar_revision_manual(
            registros[posicion],
            fila_manual,
            informe
        )

        if (
            len(informe)
            > numero_modificaciones_antes
        ):

            registros_manual += 1

    # ------------------------------------------------------
    # COMPROBACIONES FINALES
    # ------------------------------------------------------

    if (
        len(registros)
        != NUMERO_ESPERADO_REGISTROS
    ):

        raise ValueError(
            "ERROR: El número de registros "
            "ha cambiado durante el proceso."
        )

    # ------------------------------------------------------
    # ESCRITURA
    # ------------------------------------------------------

    escribir_ris(
        registros,
        ARCHIVO_RIS_SALIDA
    )

    escribir_informe(
        informe,
        ARCHIVO_INFORME
    )

    # ------------------------------------------------------
    # RECUENTOS DE CONTROL
    # ------------------------------------------------------

    con_abstract = sum(
        1
        for registro in registros
        if obtener_primer_valor(
            registro,
            "AB",
            "N2"
        )
    )

    sin_abstract = (
        len(registros)
        - con_abstract
    )

    con_doi = sum(
        1
        for registro in registros
        if obtener_primer_valor(
            registro,
            "DO"
        )
    )

    con_palabras_clave = sum(
        1
        for registro in registros
        if any(
            limpiar_espacios(valor)
            for valor in registro.get(
                "KW",
                []
            )
        )
    )

    con_idioma = sum(
        1
        for registro in registros
        if obtener_primer_valor(
            registro,
            "LA"
        )
    )

    # ------------------------------------------------------
    # RESULTADO
    # ------------------------------------------------------

    print()
    print("RESULTADO FINAL")
    print("=" * 70)

    print(
        f"Registros conservados: "
        f"{len(registros)}"
    )

    print(
        f"Registros modificados por "
        f"enriquecimiento automático: "
        f"{registros_enriquecidos}"
    )

    print(
        f"Registros modificados mediante "
        f"revisión manual: "
        f"{registros_manual}"
    )

    print(
        f"Modificaciones individuales realizadas: "
        f"{len(informe)}"
    )

    print()

    print(
        f"Registros con abstract: "
        f"{con_abstract}"
    )

    print(
        f"Registros sin abstract: "
        f"{sin_abstract}"
    )

    print(
        f"Registros con palabras clave: "
        f"{con_palabras_clave}"
    )

    print(
        f"Registros con DOI: "
        f"{con_doi}"
    )

    print(
        f"Registros con idioma identificado: "
        f"{con_idioma}"
    )

    print()

    print(
        "RIS enriquecido generado:"
    )

    print(
        ARCHIVO_RIS_SALIDA
    )

    print()

    print(
        "Informe de modificaciones:"
    )

    print(
        ARCHIVO_INFORME
    )


if __name__ == "__main__":
    principal()